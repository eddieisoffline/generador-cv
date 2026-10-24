"""DOCX export using python-docx directly (no external office application).

The layout follows the Harvard-style bullet resume template: Calibri 11 pt on
US Letter, a ruled header, centred section headings, and entry lines with the
location or dates pushed to the right margin.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from generators.formatting import date_range, display_url
from models.cv_model import (
    Award, Certification, CVProfile, Education, Experience, LeadershipActivity,
    Project, Publication, VolunteerExperience,
)
from utils.localization import t

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips
except ImportError:  # Reported when an export is requested.
    Document = None


class DocxGenerationError(RuntimeError):
    pass


# Measurements in twips, taken from the template.
PAGE_WIDTH, PAGE_HEIGHT = 12240, 15840
MARGIN_TOP, MARGIN_SIDE, MARGIN_BOTTOM = 720, 605, 280
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN_SIDE
ENTRY_INDENT = 119
BULLET_INDENT, BULLET_HANGING = 839, 360
ENTRY_GAP = 11  # points; one blank 11 pt line between entries
FONT = "Calibri"
SEPARATOR = " • "


def _join(*values: str) -> str:
    return SEPARATOR.join(value.strip() for value in values if value.strip())


def _set_font(style) -> None:
    """Pin the style to Calibri; theme fonts would otherwise win over w:ascii."""
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT)


def _configure(document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Twips(PAGE_WIDTH), Twips(PAGE_HEIGHT)
    section.top_margin, section.bottom_margin = Twips(MARGIN_TOP), Twips(MARGIN_BOTTOM)
    section.left_margin = section.right_margin = Twips(MARGIN_SIDE)
    for name in ("Normal", "Heading 1"):
        style = document.styles[name]
        _set_font(style)
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0
    heading = document.styles["Heading 1"]
    heading.font.bold = True
    heading.font.italic = False
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(ENTRY_GAP)


def _bottom_border(paragraph) -> None:
    border = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
        border.set(qn(f"w:{key}"), value)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    paragraph._p.get_or_add_pPr().insert_element_before(
        borders, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr"
    )


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    value.set(qn("xml:space"), "preserve")
    run.append(value)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    ))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _paragraph(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Twips(ENTRY_INDENT)
    return paragraph


def _entry_line(document, left: str, right: str = "", *, bold_left: bool = True) -> None:
    """Left text with the right text flushed to the right margin."""
    left, right = left.strip(), right.strip()
    if not (left or right):
        return
    paragraph = _paragraph(document)
    paragraph.add_run(left).bold = bold_left or None
    if right:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Twips(TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("\t" + right)


def _text_line(document, text: str, label: str = "") -> None:
    text, label = text.strip(), label.strip()
    if not (text or label):
        return
    paragraph = _paragraph(document)
    if label:
        paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text)


def _link_line(document, url: str) -> None:
    if url.strip():
        _add_hyperlink(_paragraph(document), display_url(url), url.strip())


def _bullet(document, text: str) -> None:
    if not text.strip():
        return
    paragraph = document.add_paragraph(text.strip(), style="List Bullet")
    paragraph.paragraph_format.left_indent = Twips(BULLET_INDENT)
    paragraph.paragraph_format.first_line_indent = Twips(-BULLET_HANGING)


def _section(document, title: str) -> None:
    document.add_paragraph(title, style="Heading 1")


def _entries(document, items: Sequence[Any], render: Callable[[Any, Any], None]) -> None:
    """Render entries with a blank line's worth of space between them."""
    for index, item in enumerate(items):
        start = len(document.paragraphs)
        render(document, item)
        if index and len(document.paragraphs) > start:
            document.paragraphs[start].paragraph_format.space_before = Pt(ENTRY_GAP)


def _header(document, profile: CVProfile) -> None:
    name = document.add_paragraph()
    _bottom_border(name)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.add_run(profile.name.strip() or "CV").bold = True
    contact = profile.contact
    values = [
        (contact.location.strip(), ""),
        (contact.email.strip(), f"mailto:{contact.email.strip()}"),
        (contact.phone.strip(), ""),
        *((display_url(url), url.strip()) for url in (
            contact.website, contact.linkedin, contact.github
        )),
    ]
    values = [(text, url) for text, url in values if text]
    if not values:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    for index, (text, url) in enumerate(values):
        if index:
            paragraph.add_run(SEPARATOR)
        if url:
            _add_hyperlink(paragraph, text, url)
        else:
            paragraph.add_run(text)


def _bullets(document, values: Iterable[str]) -> None:
    for value in values:
        _bullet(document, value)


def _education(document, item: Education) -> None:
    _entry_line(document, item.institution, item.location)
    _entry_line(document, item.degree, date_range(item.start_date, item.end_date), bold_left=False)
    for detail in item.details:
        _text_line(document, detail)


def _experience(document, item: Experience) -> None:
    _entry_line(document, item.company, item.location)
    _entry_line(document, item.role, date_range(item.start_date, item.end_date))
    _bullets(document, (bullet.text for bullet in item.bullets))


def _leadership(document, item: LeadershipActivity) -> None:
    _entry_line(document, item.organization, item.location)
    _entry_line(document, item.title, date_range(item.start_date, item.end_date))
    _text_line(document, item.description)
    _bullets(document, (bullet.text for bullet in item.bullets))


def _volunteer(document, item: VolunteerExperience) -> None:
    _entry_line(document, item.organization, item.location)
    _entry_line(document, item.role, date_range(item.start_date, item.end_date))
    _text_line(document, item.description)
    _bullets(document, (bullet.text for bullet in item.bullets))


def _project(document, item: Project) -> None:
    _entry_line(document, item.name)
    _text_line(document, item.description)
    _text_line(document, ", ".join(item.technologies))
    _link_line(document, item.url)
    _bullets(document, (bullet.text for bullet in item.bullets))


def _certification(document, item: Certification) -> None:
    _entry_line(document, item.name, date_range(item.date, item.expiry_date))
    _text_line(document, _join(item.issuer, item.credential_id))
    _text_line(document, item.description)
    _link_line(document, item.url)
    _bullets(document, item.details)


def _award(document, item: Award) -> None:
    _entry_line(document, item.name, item.date)
    _text_line(document, item.issuer)
    _text_line(document, item.description)
    _link_line(document, item.url)
    _bullets(document, item.details)


def _publication(document, item: Publication) -> None:
    _entry_line(document, item.title, item.date)
    _text_line(document, _join(item.authors, item.publisher))
    _text_line(document, item.description)
    _link_line(document, item.url)


def _skills_interests(document, profile: CVProfile) -> None:
    lines = []
    if profile.section_is_active("skills"):
        lines += [(group.category, ", ".join(group.items))
                  for group in profile.skills if group.category or group.items]
    extra = []
    if profile.languages and profile.section_is_active("languages"):
        extra.append((t(profile.locale, "languages"), ", ".join(profile.languages)))
    if profile.interests and profile.section_is_active("interests"):
        extra.append((t(profile.locale, "interests"), ", ".join(profile.interests)))
    if not lines and not extra:
        return
    _section(document, t(profile.locale, "skills_interests" if extra else "skills"))
    for label, text in lines + extra:
        _text_line(document, text, label)


def _document(profile: CVProfile):
    if Document is None:
        raise DocxGenerationError("Install python-docx to enable Word export.")
    document = Document()
    document.core_properties.title = profile.name
    document.core_properties.author = profile.name
    _configure(document)
    _header(document, profile)

    def entries(key: str, items: Sequence[Any], render: Callable[[Any, Any], None]) -> None:
        if items and profile.section_is_active(key):
            _section(document, t(profile.locale, key))
            _entries(document, items, render)

    # Section order follows the template; title and summary are not exported.
    entries("education", profile.education, _education)
    entries("experience", profile.experience, _experience)
    entries("leadership_activities", profile.leadership_activities, _leadership)
    _skills_interests(document, profile)
    entries("projects", profile.projects, _project)
    entries("volunteer_experience", profile.volunteer_experience, _volunteer)
    entries("certifications", profile.certifications, _certification)
    entries("awards", profile.awards, _award)
    entries("publications", profile.publications, _publication)
    return document


def generate_docx(profile: CVProfile, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _document(profile).save(str(output_path))
    return output_path


def generate_docx_bytes(profile: CVProfile) -> bytes:
    stream = BytesIO()
    _document(profile).save(stream)
    return stream.getvalue()


class DocxGenerator:
    """Object-oriented facade for direct python-docx generation."""

    def generate(self, profile: CVProfile, output_path: Path) -> Path:
        return generate_docx(profile, output_path)

    def generate_bytes(self, profile: CVProfile) -> bytes:
        return generate_docx_bytes(profile)
